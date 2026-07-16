from io import BytesIO

from httpx import AsyncClient
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

TEMPLATE_CODIGO = "VINCULO"


def _xlsx_vinculo_minimo() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.cell(row=3, column=column_index_from_string("A"), value="01/01/2020")
    ws.cell(row=3, column=column_index_from_string("C"), value="V001")
    ws.cell(row=3, column=column_index_from_string("J"), value="12345678900")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def test_vinculo_dicionario_parcial_expoe_apenas_os_campos_validados(client: AsyncClient) -> None:
    response = await client.get(f"/templates/{TEMPLATE_CODIGO}")
    assert response.status_code == 200
    campos = {c["campo"] for c in response.json()["campos"]}
    assert campos == {"DTADMISSAOVINC", "CDMATRICULA", "NRCPFPESSOA", "NRVINCULOM"}


async def test_vinculo_preview_funciona_com_dicionario_parcial(client: AsyncClient) -> None:
    response = await client.post(
        f"/templates/{TEMPLATE_CODIGO}/preview",
        files={"arquivo": ("vinculo.xlsx", _xlsx_vinculo_minimo(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"nr_org": "3260"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total_linhas"] == 1
    assert body["validos"] == 1


async def test_vinculo_gerar_script_falha_pois_nao_ha_script_cadastrado(client: AsyncClient) -> None:
    """Vínculo tem dicionário parcial e nenhum TemplateScript cadastrado (Seção 19) — a
    geração de script deve falhar de forma explícita, não silenciosa."""
    response = await client.post(
        f"/templates/{TEMPLATE_CODIGO}/gerar-script",
        files={"arquivo": ("vinculo.xlsx", _xlsx_vinculo_minimo(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"nr_org": "3260"},
    )
    assert response.status_code == 400
